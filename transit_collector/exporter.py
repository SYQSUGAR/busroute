from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from transit_collector.db import TransitDB
from transit_collector.gis import GISProcessor


class Exporter:
    def __init__(self, db: TransitDB):
        self.db = db

    def export_all(self, out_dir: str | Path) -> dict[str, str]:
        """Export relational tables plus standards-compliant GIS layers."""
        out = Path(out_dir)
        out.mkdir(parents=True, exist_ok=True)

        files: dict[str, str] = {}
        files["lines_csv"] = str(self._export_lines_csv(out / "bus_lines.csv"))
        files["line_stops_csv"] = str(self._export_line_stops_csv(out / "bus_line_stops.csv"))
        files["stations_csv"] = str(self._export_stations_csv(out / "stations_merged.csv"))
        files["station_lines_csv"] = str(self._export_station_lines_csv(out / "station_lines.csv"))
        files["raw_stops_csv"] = str(self._export_raw_stops_csv(out / "stops_raw.csv"))
        files["xlsx"] = str(self._export_xlsx(out / "公交线路与站点.xlsx"))

        gis_files = GISProcessor(self.db).export_base_layers(out, write_shp=True)
        files.update(gis_files)

        files["summary"] = str(self._export_summary(out / "summary.json"))
        return files

    def _write_query_csv(self, path: Path, query: str):
        cur = self.db.conn.execute(query)
        headers = [d[0] for d in cur.description]
        with path.open("w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(headers)
            for row in cur:
                writer.writerow([row[h] for h in headers])
        return path

    def _export_lines_csv(self, path):
        return self._write_query_csv(path, """
            SELECT line_id,name,type,citycode,crawl_city_adcode,start_stop,end_stop,
                   start_time,end_time,distance,loop,status,direc,company,basic_price,total_price,polyline
            FROM bus_lines ORDER BY crawl_city_adcode,name,line_id
        """)

    def _export_line_stops_csv(self, path):
        return self._write_query_csv(path, """
            SELECT ls.line_id, bl.name AS line_name, ls.sequence,
                   COALESCE(sm.station_id,'') AS station_id,
                   ls.stop_id AS raw_stop_id, ls.stop_name,
                   ls.longitude,ls.latitude
            FROM line_stops ls
            JOIN bus_lines bl ON bl.line_id=ls.line_id
            LEFT JOIN station_members sm ON sm.stop_id=ls.stop_id
            ORDER BY bl.name,ls.line_id,ls.sequence
        """)

    def _export_stations_csv(self, path):
        return self._write_query_csv(path, """
            SELECT sg.station_id,sg.name,sg.longitude,sg.latitude,sg.member_count,sg.line_count,
                   GROUP_CONCAT(DISTINCT sl.line_id) AS line_ids,
                   GROUP_CONCAT(DISTINCT bl.name) AS line_names,
                   GROUP_CONCAT(DISTINCT sm.stop_id) AS raw_stop_ids
            FROM station_groups sg
            LEFT JOIN station_members sm ON sm.station_id=sg.station_id
            LEFT JOIN station_lines sl ON sl.station_id=sg.station_id
            LEFT JOIN bus_lines bl ON bl.line_id=sl.line_id
            GROUP BY sg.station_id
            ORDER BY sg.name,sg.station_id
        """)

    def _export_station_lines_csv(self, path):
        return self._write_query_csv(path, """
            SELECT sl.station_id,sg.name AS station_name,sl.line_id,bl.name AS line_name,
                   bl.start_stop,bl.end_stop
            FROM station_lines sl
            JOIN station_groups sg ON sg.station_id=sl.station_id
            JOIN bus_lines bl ON bl.line_id=sl.line_id
            ORDER BY sg.name,bl.name
        """)

    def _export_raw_stops_csv(self, path):
        return self._write_query_csv(path, """
            SELECT stop_id,name,norm_name,longitude,latitude,adcode,citycode,source
            FROM raw_stops ORDER BY name,stop_id
        """)

    def _sheet_from_query(self, wb: Workbook, title: str, query: str):
        ws = wb.create_sheet(title)
        cur = self.db.conn.execute(query)
        headers = [d[0] for d in cur.description]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in cur:
            ws.append([row[h] for h in headers])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, h in enumerate(headers, start=1):
            width = min(max(12, len(str(h)) * 2 + 2), 40)
            ws.column_dimensions[get_column_letter(i)].width = width
        return ws

    def _export_xlsx(self, path: Path):
        wb = Workbook()
        wb.remove(wb.active)
        self._sheet_from_query(wb, "公交线路", """
            SELECT line_id,name,type,start_stop,end_stop,start_time,end_time,distance,loop,status,
                   direc,company,basic_price,total_price,crawl_city_adcode
            FROM bus_lines ORDER BY crawl_city_adcode,name,line_id
        """)
        self._sheet_from_query(wb, "线路-站点序列", """
            SELECT ls.line_id,bl.name AS line_name,ls.sequence,
                   COALESCE(sm.station_id,'') AS station_id,ls.stop_id AS raw_stop_id,
                   ls.stop_name,ls.longitude,ls.latitude
            FROM line_stops ls
            JOIN bus_lines bl ON bl.line_id=ls.line_id
            LEFT JOIN station_members sm ON sm.stop_id=ls.stop_id
            ORDER BY bl.name,ls.line_id,ls.sequence
        """)
        self._sheet_from_query(wb, "合并站点", """
            SELECT sg.station_id,sg.name,sg.longitude,sg.latitude,sg.member_count,sg.line_count,
                   GROUP_CONCAT(DISTINCT bl.name) AS line_names,
                   GROUP_CONCAT(DISTINCT sl.line_id) AS line_ids
            FROM station_groups sg
            LEFT JOIN station_lines sl ON sl.station_id=sg.station_id
            LEFT JOIN bus_lines bl ON bl.line_id=sl.line_id
            GROUP BY sg.station_id ORDER BY sg.name
        """)
        self._sheet_from_query(wb, "站点-线路关系", """
            SELECT sl.station_id,sg.name AS station_name,sl.line_id,bl.name AS line_name,
                   bl.start_stop,bl.end_stop
            FROM station_lines sl
            JOIN station_groups sg ON sg.station_id=sl.station_id
            JOIN bus_lines bl ON bl.line_id=sl.line_id
            ORDER BY sg.name,bl.name
        """)
        wb.save(path)
        return path

    def _export_summary(self, path: Path):
        counts = self.db.counts()
        summary = {
            "counts": counts,
            "scope": self.db.get_meta("scope", ""),
            "completed": self.db.get_meta("completed", "0"),
            "raw_coordinate_system": "AMap/GCJ-02",
            "gis_coordinate_reference_system": (
                "GIS outputs use EPSG:4326 after approximate iterative GCJ-02 -> WGS84 conversion"
            ),
            "primary_spatial_format": "GeoPackage",
            "notes": [
                "stations_merged 为按站名 + 空间距离聚合后的站点。",
                "bus_line_stops 保留每条线路的站点序号与原始坐标。",
                "transit_gis.gpkg 包含 bus_stations、bus_routes、route_stops 标准矢量图层。",
                "shp/ 目录提供 ArcGIS 兼容 Shapefile；GeoPackage 应作为优先交换格式。",
                "高德 POI 搜索与公交高级 API 有配额和检索上限，正式成果建议结合主管部门台账抽检。",
            ],
        }
        path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), "utf-8")
        return path
